import streamlit as st
import pandas as pd
import random
import re
import io
import zipfile

from docx import Document


# ============================================================
# CONFIG
# ============================================================

st.set_page_config(
    page_title="Exam Generator",
    page_icon="📝",
    layout="wide"
)


# ============================================================
# WORD PARSER
# ============================================================

def read_questions_from_bytes(file_bytes):

    doc = Document(
        io.BytesIO(file_bytes)
    )

    questions = []

    current = None

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        # ====================================================
        # QUESTION HEADER
        #
        # Câu 01 – QB001
        # [Khái niệm cơ bản] Nội dung câu hỏi...
        # ====================================================

        match = re.match(
            r"^Câu\s*(\d+)\s*[-–—]\s*(QB\d+)\s*(?:\n)?(.*)",
            text,
            re.IGNORECASE | re.DOTALL
        )

        if match:

            # Save previous question
            if current:
                questions.append(current)

            question_no = int(
                match.group(1)
            )

            question_id = (
                match.group(2)
                .upper()
                .strip()
            )

            question = (
                match.group(3)
                .strip()
            )

            # Remove [....]
            question = re.sub(
                r"^\[.*?\]\s*",
                "",
                question
            )

            current = {

                "Question_No":
                    question_no,

                "Question_ID":
                    question_id,

                "Question":
                    question,

                "A": "",
                "B": "",
                "C": "",
                "D": "",

                "Correct_Answer":
                    "",

                "CLO":
                    "",

                "Bloom":
                    ""
            }

            continue

        # ====================================================
        # ANSWER OPTIONS
        # ====================================================

        if current:

            option_match = re.match(
                r"^([ABCD])[\.\:\)]\s*(.*)",
                text,
                re.IGNORECASE
            )

            if option_match:

                letter = (
                    option_match
                    .group(1)
                    .upper()
                )

                current[letter] = (
                    option_match
                    .group(2)
                    .strip()
                )

                continue

        # ====================================================
        # METADATA
        #
        # Đáp án: C | CLO: CLO1 | Bloom: Nhận biết
        #
        # Correct_Answer: C | CLO: CLO1 | Bloom: Nhận biết
        # ====================================================

        if current:

            answer_match = re.search(
                r"(?:Đáp án|Correct_Answer)\s*:\s*([ABCD])",
                text,
                re.IGNORECASE
            )

            clo_match = re.search(
                r"CLO\s*:\s*(CLO\d+)",
                text,
                re.IGNORECASE
            )

            bloom_match = re.search(
                r"Bloom\s*:\s*(.+)",
                text,
                re.IGNORECASE
            )

            if answer_match:

                current["Correct_Answer"] = (
                    answer_match
                    .group(1)
                    .upper()
                )

            if clo_match:

                current["CLO"] = (
                    clo_match
                    .group(1)
                    .upper()
                )

            if bloom_match:

                current["Bloom"] = (
                    bloom_match
                    .group(1)
                    .strip()
                )

    # Save final question

    if current:
        questions.append(current)

    return questions


# ============================================================
# CREATE WORD TEMPLATE
# ============================================================

def create_word_template(
    number_of_questions
):

    doc = Document()

    # Title

    title = doc.add_heading(
        "QUESTION BANK",
        level=1
    )

    title.alignment = 1

    doc.add_paragraph(
        "Template ngân hàng câu hỏi"
    )

    doc.add_paragraph(
        "Giữ nguyên cấu trúc Question_ID, "
        "CLO và Bloom khi nhập dữ liệu."
    )

    doc.add_paragraph("")

    # Questions

    for i in range(
        1,
        number_of_questions + 1
    ):

        question_id = (
            f"QB{i:03d}"
        )

        # Header + question

        p = doc.add_paragraph()

        run = p.add_run(
            f"Câu {i:02d} – {question_id}\n"
        )

        run.bold = True

        p.add_run(
            "[Chủ đề] "
            "Nhập nội dung câu hỏi tại đây."
        )

        # Options

        doc.add_paragraph(
            "A. Nhập phương án A."
        )

        doc.add_paragraph(
            "B. Nhập phương án B."
        )

        doc.add_paragraph(
            "C. Nhập phương án C."
        )

        doc.add_paragraph(
            "D. Nhập phương án D."
        )

        # Metadata

        doc.add_paragraph(
            "Đáp án: A | CLO: CLO1 | "
            "Bloom: Nhận biết"
        )

        doc.add_paragraph("")

    buffer = io.BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer.getvalue()


# ============================================================
# CREATE EXAM WORD
# ============================================================

def create_exam_word(
    exam_df,
    exam_id
):

    doc = Document()

    title = doc.add_heading(
        f"ĐỀ THI {exam_id}",
        level=1
    )

    title.alignment = 1

    answer_rows = []

    for question_no, (_, row) in enumerate(
        exam_df.iterrows(),
        start=1
    ):

        # ====================================================
        # QUESTION
        # ====================================================

        doc.add_paragraph(
            f"Câu {question_no}. "
            f"{row['Question']}"
        )

        # ====================================================
        # OPTIONS
        # ====================================================

        doc.add_paragraph(
            f"A. {row['A']}"
        )

        doc.add_paragraph(
            f"B. {row['B']}"
        )

        doc.add_paragraph(
            f"C. {row['C']}"
        )

        doc.add_paragraph(
            f"D. {row['D']}"
        )

        # ====================================================
        # ANSWER DATABASE
        # ====================================================

        answer_rows.append({

            "Exam_ID":
                exam_id,

            "Question_No":
                question_no,

            "Question_ID":
                row["Question_ID"],

            "Answer":
                row["Correct_Answer"],

            "CLO":
                row["CLO"],

            "Bloom":
                row["Bloom"]
        })

    buffer = io.BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return (
        buffer.getvalue(),
        pd.DataFrame(answer_rows)
    )


# ============================================================
# GENERATE ONE EXAM
# ============================================================

def generate_exam(
    data,
    number_of_questions,
    clo_percent
):

    selected = []

    # ========================================================
    # SELECT EACH CLO
    # ========================================================

    for clo, percent in clo_percent.items():

        required = round(
            number_of_questions * percent
        )

        if required <= 0:
            continue

        pool = data[
            data["CLO"] == clo
        ]

        if len(pool) < required:

            raise ValueError(
                f"{clo}: cần {required} câu "
                f"nhưng chỉ có {len(pool)} câu."
            )

        sample = pool.sample(
            n=required,
            replace=False
        )

        selected.append(sample)

    if not selected:

        raise ValueError(
            "Không có câu hỏi được chọn."
        )

    exam = pd.concat(
        selected,
        ignore_index=True
    )

    # ========================================================
    # FIX ROUNDING
    # ========================================================

    if len(exam) < number_of_questions:

        used_ids = set(
            exam["Question_ID"]
        )

        remaining = data[
            ~data["Question_ID"]
            .isin(used_ids)
        ]

        need = (
            number_of_questions
            - len(exam)
        )

        if len(remaining) < need:

            raise ValueError(
                "Không đủ câu để bổ sung."
            )

        extra = remaining.sample(
            n=need,
            replace=False
        )

        exam = pd.concat(
            [
                exam,
                extra
            ],
            ignore_index=True
        )

    elif len(exam) > number_of_questions:

        exam = exam.sample(
            n=number_of_questions,
            replace=False
        )

    # ========================================================
    # RANDOM QUESTION ORDER
    # ========================================================

    exam = exam.sample(
        frac=1
    ).reset_index(
        drop=True
    )

    return exam


# ============================================================
# CREATE ANSWER EXCEL
# ============================================================

def create_answer_excel(
    answer_database
):

    output = io.BytesIO()

    # ========================================================
    # ANSWER KEY
    # ========================================================

    answer_key = answer_database.pivot(
        index="Exam_ID",
        columns="Question_No",
        values="Answer"
    )

    answer_key.columns = [
        f"Câu {c}"
        for c in answer_key.columns
    ]

    answer_key = (
        answer_key
        .reset_index()
    )

    # ========================================================
    # DATABASE
    # ========================================================

    question_database = (
        answer_database[
            [
                "Exam_ID",
                "Question_No",
                "Question_ID",
                "Answer",
                "CLO",
                "Bloom"
            ]
        ]
    )

    # ========================================================
    # WRITE EXCEL
    # ========================================================

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        question_database.to_excel(
            writer,
            sheet_name="Exam_Answers",
            index=False
        )

        answer_key.to_excel(
            writer,
            sheet_name="Answer_Key",
            index=False
        )

    output.seek(0)

    return output.getvalue()


# ============================================================
# MAIN APP
# ============================================================

st.title(
    "📝 Question Bank & Exam Generator"
)

st.caption(
    "Tạo ngân hàng câu hỏi và trộn đề thi "
    "theo CLO."
)


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.header(
    "⚙️ Chức năng"
)

mode = st.sidebar.radio(
    "Chọn chức năng",
    [
        "Tạo Word Template",
        "Tạo đề thi"
    ]
)


# ============================================================
# MODE 1: TEMPLATE
# ============================================================

if mode == "Tạo Word Template":

    st.header(
        "📄 Tạo Word Template"
    )

    st.write(
        "Tạo file Word có cấu trúc chuẩn "
        "để nhập ngân hàng câu hỏi."
    )

    number_of_questions = st.number_input(
        "Số câu",
        min_value=1,
        max_value=5000,
        value=50,
        step=10
    )

    if st.button(
        "📄 TẠO WORD TEMPLATE",
        type="primary"
    ):

        file_bytes = create_word_template(
            number_of_questions
        )

        st.success(
            f"✓ Đã tạo template {number_of_questions} câu."
        )

        st.download_button(
            label="⬇️ Download Word Template",
            data=file_bytes,
            file_name=(
                f"Question_Bank_Template_"
                f"{number_of_questions}_Cau.docx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
            use_container_width=True
        )


# ============================================================
# MODE 2: EXAM GENERATOR
# ============================================================

else:

    st.header(
        "📝 Tạo đề thi"
    )

    uploaded_file = st.file_uploader(
        "Upload Question Bank Word",
        type=["docx"]
    )

    if uploaded_file:

        # ====================================================
        # READ WORD
        # ====================================================

        try:

            questions = (
                read_questions_from_bytes(
                    uploaded_file.getvalue()
                )
            )

        except Exception as e:

            st.error(
                f"Lỗi đọc file Word: {e}"
            )

            st.stop()

        data = pd.DataFrame(
            questions
        )

        # ====================================================
        # BASIC CHECK
        # ====================================================

        if len(data) == 0:

            st.error(
                "❌ Không đọc được câu hỏi nào."
            )

            st.stop()

        st.success(
            f"✓ Đọc được {len(data)} câu hỏi."
        )

        # ====================================================
        # METRICS
        # ====================================================

        c1, c2, c3 = st.columns(3)

        with c1:

            st.metric(
                "Số câu",
                len(data)
            )

        with c2:

            st.metric(
                "Số CLO",
                data["CLO"].nunique()
            )

        with c3:

            st.metric(
                "Số Bloom",
                data["Bloom"].nunique()
            )

        # ====================================================
        # VALIDATE
        # ====================================================

        required_columns = [
            "Question_ID",
            "Question",
            "A",
            "B",
            "C",
            "D",
            "Correct_Answer",
            "CLO",
            "Bloom"
        ]

        missing = []

        for column in required_columns:

            if column not in data.columns:

                missing.append(column)

        if missing:

            st.error(
                "Thiếu dữ liệu: "
                + ", ".join(missing)
            )

            st.stop()

        # ====================================================
        # DUPLICATE ID
        # ====================================================

        duplicates = (
            data[
                data["Question_ID"]
                .duplicated(
                    keep=False
                )
            ]
        )

        if len(duplicates) > 0:

            st.error(
                "❌ Question_ID bị trùng:"
            )

            st.dataframe(
                duplicates[
                    [
                        "Question_No",
                        "Question_ID"
                    ]
                ],
                use_container_width=True
            )

            st.stop()

        # ====================================================
        # CLO
        # ====================================================

        st.subheader(
            "🎯 Phân bố CLO"
        )

        clo_list = sorted(
            data["CLO"]
            .dropna()
            .unique()
            .tolist()
        )

        # ====================================================
        # EXAM CONFIG
        # ====================================================

        col1, col2 = st.columns(2)

        with col1:

            number_of_questions = (
                st.number_input(
                    "Số câu mỗi đề",
                    min_value=1,
                    max_value=len(data),
                    value=min(
                        20,
                        len(data)
                    ),
                    step=1
                )
            )

        with col2:

            number_of_exams = (
                st.number_input(
                    "Số đề",
                    min_value=1,
                    max_value=500,
                    value=10,
                    step=1
                )
            )

        # ====================================================
        # CLO PERCENT
        # ====================================================

        clo_percent = {}

        columns = st.columns(
            min(
                4,
                len(clo_list)
            )
        )

        for i, clo in enumerate(
            clo_list
        ):

            count = len(
                data[
                    data["CLO"] == clo
                ]
            )

            with columns[
                i % len(columns)
            ]:

                percent = st.number_input(
                    f"{clo} (%)",
                    min_value=0.0,
                    max_value=100.0,
                    value=(
                        100.0
                        / len(clo_list)
                    ),
                    step=5.0,
                    key=f"clo_{clo}"
                )

                clo_percent[clo] = (
                    percent / 100
                )

                st.caption(
                    f"Ngân hàng: {count} câu"
                )

        # ====================================================
        # TOTAL
        # ====================================================

        total = (
            sum(clo_percent.values())
            * 100
        )

        if abs(total - 100) > 0.01:

            st.warning(
                f"⚠️ Tổng CLO = {total:.1f}% "
                "→ phải bằng 100%."
            )

        else:

            st.success(
                "✓ Tổng CLO = 100%"
            )

        # ====================================================
        # PREVIEW
        # ====================================================

        with st.expander(
            "👁 Xem ngân hàng câu hỏi"
        ):

            st.dataframe(
                data,
                use_container_width=True,
                hide_index=True
            )

        # ====================================================
        # GENERATE
        # ====================================================

        st.divider()

        if st.button(
            "🚀 TẠO ĐỀ THI",
            type="primary",
            use_container_width=True
        ):

            if abs(total - 100) > 0.01:

                st.error(
                    "Tổng tỷ lệ CLO phải bằng 100%."
                )

                st.stop()

            # =================================================
            # CAPACITY CHECK
            # =================================================

            errors = []

            for clo, percent in (
                clo_percent.items()
            ):

                required = round(
                    number_of_questions
                    * percent
                )

                available = len(
                    data[
                        data["CLO"] == clo
                    ]
                )

                if required > available:

                    errors.append(
                        f"{clo}: cần "
                        f"{required} câu, "
                        f"chỉ có "
                        f"{available} câu."
                    )

            if errors:

                for error in errors:

                    st.error(
                        error
                    )

                st.stop()

            # =================================================
            # GENERATE
            # =================================================

            all_answers = []

            files = {}

            progress = st.progress(0)

            status = st.empty()

            for i in range(
                1,
                number_of_exams + 1
            ):

                exam_id = (
                    f"DE{i:02d}"
                )

                status.write(
                    f"Đang tạo {exam_id}..."
                )

                exam = generate_exam(
                    data,
                    number_of_questions,
                    clo_percent
                )

                word_bytes, answers = (
                    create_exam_word(
                        exam,
                        exam_id
                    )
                )

                files[
                    f"{exam_id}.docx"
                ] = word_bytes

                all_answers.append(
                    answers
                )

                progress.progress(
                    i / number_of_exams
                )

            # =================================================
            # ANSWER DATABASE
            # =================================================

            answer_database = pd.concat(
                all_answers,
                ignore_index=True
            )

            # =================================================
            # EXCEL
            # =================================================

            excel_bytes = (
                create_answer_excel(
                    answer_database
                )
            )

            files[
                "Answer_Key.xlsx"
            ] = excel_bytes

            # =================================================
            # ZIP
            # =================================================

            zip_buffer = io.BytesIO()

            with zipfile.ZipFile(
                zip_buffer,
                "w",
                zipfile.ZIP_DEFLATED
            ) as z:

                for filename, content in (
                    files.items()
                ):

                    z.writestr(
                        filename,
                        content
                    )

            zip_buffer.seek(0)

            status.success(
                "✓ Hoàn thành!"
            )

            st.success(
                f"🎉 Đã tạo "
                f"{number_of_exams} đề, "
                f"mỗi đề "
                f"{number_of_questions} câu."
            )

            # =================================================
            # DOWNLOAD
            # =================================================

            st.download_button(
                label="📦 DOWNLOAD ZIP",
                data=zip_buffer.getvalue(),
                file_name=(
                    "Exam_Generator_Output.zip"
                ),
                mime="application/zip",
                use_container_width=True
            )